"""Bundle the supplied 30-topic collocations workbook into the Site app."""
import json
import re
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
DOCX = ROOT / 'content/max-speaking-workbook.docx'
APP = ROOT / 'dist/app.js'
document = Document(DOCX)
headings = [p.text.strip() for p in document.paragraphs if p.style.name == 'Heading 1']
if len(headings) != 30 or len(document.tables) != 1500:
    raise ValueError('Workbook must contain 30 topics and 1,500 entries')

prefixes = ['Uzbek: ', 'Example 1: ', 'Example 2: ', 'Siz tarjima qiling 1: ', 'Siz tarjima qiling 2: ']
topics = []
for topic_index, heading in enumerate(headings):
    match = re.fullmatch(r'(\d+)\. (.+)', heading)
    if not match or int(match.group(1)) != topic_index + 1:
        raise ValueError('Unexpected topic heading: ' + heading)
    entries = []
    for entry_index, table in enumerate(document.tables[topic_index * 50:(topic_index + 1) * 50]):
        lines = [p.text.strip() for p in table.cell(0, 0).paragraphs]
        if len(lines) != 6:
            raise ValueError(f'Entry {topic_index + 1}.{entry_index + 1} has {len(lines)} lines')
        label = re.fullmatch(r'(\d+)\. (.+?)\s+\[(Speaking|Writing)\]', lines[0])
        if not label or int(label.group(1)) != entry_index + 1:
            raise ValueError('Unexpected collocation label: ' + lines[0])
        values = []
        for line, prefix in zip(lines[1:], prefixes):
            if not line.startswith(prefix) or not line[len(prefix):].strip():
                raise ValueError('Missing ' + prefix + ' in ' + lines[0])
            values.append(line[len(prefix):].strip())
        if label.group(3) != ('Speaking' if entry_index < 30 else 'Writing'):
            raise ValueError('Unexpected Speaking/Writing classification')
        entries.append(dict(id=f't{topic_index + 1:02}e{entry_index + 1:02}', term=label.group(2), kind=label.group(3), uz=values[0], examples=values[1:3], prompts=values[3:5]))
    topics.append(dict(id=f't{topic_index + 1:02}', title=match.group(2), entries=entries))

data = json.dumps({'topics': topics}, ensure_ascii=False, separators=(',', ':'))
start = '// MAX_SPEAKING_DATA_START'
end = '// MAX_SPEAKING_DATA_END'
app = APP.read_text(encoding='utf-8')
if app.count(start) != 1 or app.count(end) != 1:
    raise ValueError('MAX SPEAKING data markers missing from app.js')
app = re.sub(re.escape(start) + r'.*?' + re.escape(end), lambda _: f'{start}\nconst maxSpeakingBank={data};\n{end}', app, count=1, flags=re.S)
APP.write_text(app, encoding='utf-8')
print(f'MAX SPEAKING built: {len(topics)} topics, {sum(len(t["entries"]) for t in topics)} collocations')
